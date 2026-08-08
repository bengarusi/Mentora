"""Real-file extraction: each format is generated for the test, then read back
through the same extractor the upload endpoint uses."""

import io

import pytest

from app.files.extraction import (
    ExtractionError,
    ExtractorUnavailable,
    get_extractor,
)


def _make_pdf(text: str) -> bytes:
    """A minimal but structurally valid PDF (correct xref table and EOF) with a
    real text layer — pypdf rejects hand-waved PDFs, as it should."""
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        None,
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs[3] = (
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream"
    )

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj".encode() + body + b"endobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer<</Size {len(objs) + 1}/Root 1 0 R>>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_pdf_with_a_text_layer_is_extracted():
    extractor = get_extractor("application/pdf", "notes.pdf")
    result = extractor.extract(_make_pdf("Rounding to the nearest ten"), "notes.pdf")
    assert "Rounding to the nearest ten" in result.text
    assert result.page_count == 1


def test_scanned_pdf_is_parked_for_ocr_rather_than_failed():
    """A page with no text layer isn't a broken file — it needs OCR, so it must
    raise ExtractorUnavailable (parked) not ExtractionError (failed)."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    extractor = get_extractor("application/pdf", "scan.pdf")
    with pytest.raises(ExtractorUnavailable):
        extractor.extract(buf.getvalue(), "scan.pdf")


def test_corrupt_pdf_raises_a_clean_extraction_error():
    extractor = get_extractor("application/pdf", "broken.pdf")
    with pytest.raises(ExtractionError):
        extractor.extract(b"%PDF-1.4 this is not really a pdf", "broken.pdf")


def test_docx_paragraphs_and_tables_are_extracted():
    import docx

    document = docx.Document()
    document.add_paragraph("Equivalent fractions have the same value.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "1/2"
    table.rows[0].cells[1].text = "2/4"
    buf = io.BytesIO()
    document.save(buf)

    extractor = get_extractor(None, "notes.docx")
    result = extractor.extract(buf.getvalue(), "notes.docx")
    assert "Equivalent fractions have the same value." in result.text
    assert "1/2 | 2/4" in result.text  # table rows survive as readable text


def test_pptx_slides_are_extracted_with_boundaries():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Adding Fractions"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Keep the denominator the same."
    buf = io.BytesIO()
    presentation.save(buf)

    extractor = get_extractor(None, "deck.pptx")
    result = extractor.extract(buf.getvalue(), "deck.pptx")
    assert "Slide 1:" in result.text
    assert "Adding Fractions" in result.text
    assert result.page_count == 1


def test_plain_text_handles_non_utf8_encodings():
    extractor = get_extractor("text/plain", "notes.txt")
    assert "café" in extractor.extract("café".encode("utf-16"), "notes.txt").text


@pytest.mark.parametrize(
    "filename,expected",
    [
        ("a.pdf", "PdfExtractor"),
        ("a.docx", "DocxExtractor"),
        ("a.pptx", "PptxExtractor"),
        ("a.md", "PlainTextExtractor"),
        ("a.png", "ImageVisionExtractor"),
        ("a.JPG", "ImageVisionExtractor"),  # case-insensitive
    ],
)
def test_extractor_is_chosen_by_extension(filename, expected):
    extractor = get_extractor(None, filename)
    assert type(extractor).__name__ == expected


def test_unknown_type_has_no_extractor():
    assert get_extractor("application/octet-stream", "mystery.xyz") is None
