from __future__ import annotations

import base64
import io
import logging
import time
from abc import ABC, abstractmethod
from dataclasses import dataclass

from app.core.config import settings

log = logging.getLogger("app.files.extraction")


class ExtractionError(Exception):
    """Raised when an extractor was the right one but failed on this file."""


class ExtractorUnavailable(Exception):
    """Raised when an extractor's optional dependency (or API key) is missing.

    Distinct from ExtractionError: the file is a supported type, the environment
    just cannot process it yet, so the material is parked rather than failed."""


@dataclass
class ExtractedDocument:
    """Plain-text view of an uploaded file, plus whatever structure survived."""

    text: str
    page_count: int | None = None


class DocumentExtractor(ABC):
    """Turns raw upload bytes into plain text.

    One extractor per family of formats. `supports` decides ownership by MIME
    type or file extension; `extract` does the work. Adding OCR for scanned PDFs
    or a new format means adding an extractor and registering it — nothing else
    in the pipeline changes."""

    #: MIME types this extractor claims.
    content_types: tuple[str, ...] = ()
    #: Lowercase file extensions (with dot) this extractor claims.
    extensions: tuple[str, ...] = ()

    def supports(self, content_type: str | None, filename: str) -> bool:
        if content_type and content_type.split(";")[0].strip() in self.content_types:
            return True
        lowered = filename.lower()
        return any(lowered.endswith(ext) for ext in self.extensions)

    @abstractmethod
    def extract(self, data: bytes, filename: str) -> ExtractedDocument:
        """Return the document's text. Raises ExtractionError on a real failure
        and ExtractorUnavailable when a dependency is missing."""


class PlainTextExtractor(DocumentExtractor):
    """.txt / .md — decoded directly, no dependencies."""

    content_types = ("text/plain", "text/markdown", "text/x-markdown")
    extensions = (".txt", ".md", ".markdown", ".text")

    def extract(self, data: bytes, filename: str) -> ExtractedDocument:
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return ExtractedDocument(text=data.decode(encoding))
            except UnicodeDecodeError:
                continue
        raise ExtractionError("Could not decode the text file.")


class PdfExtractor(DocumentExtractor):
    """PDF via pypdf. Text-layer only — a scanned PDF yields little or nothing,
    which the service reports as UNSUPPORTED so OCR can pick it up later."""

    content_types = ("application/pdf",)
    extensions = (".pdf",)

    def extract(self, data: bytes, filename: str) -> ExtractedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ExtractorUnavailable("pypdf is not installed") from exc

        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [(page.extract_text() or "") for page in reader.pages]
        except Exception as exc:  # noqa: BLE001 - normalize parser errors
            raise ExtractionError(f"Could not read the PDF: {exc}") from exc

        text = "\n\n".join(p.strip() for p in pages if p.strip())
        if not text.strip():
            raise ExtractorUnavailable(
                "This PDF has no selectable text (it looks scanned). "
                "Image-based PDFs need OCR, which isn't enabled yet."
            )
        return ExtractedDocument(text=text, page_count=len(pages))


class DocxExtractor(DocumentExtractor):
    """Word .docx via python-docx — paragraphs plus table cells."""

    content_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    extensions = (".docx",)

    def extract(self, data: bytes, filename: str) -> ExtractedDocument:
        try:
            import docx
        except ImportError as exc:
            raise ExtractorUnavailable("python-docx is not installed") from exc

        try:
            document = docx.Document(io.BytesIO(data))
            parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        except Exception as exc:  # noqa: BLE001 - normalize parser errors
            raise ExtractionError(f"Could not read the Word file: {exc}") from exc

        return ExtractedDocument(text="\n\n".join(parts))


class PptxExtractor(DocumentExtractor):
    """PowerPoint .pptx via python-pptx — one text block per slide, so slide
    boundaries survive into chunking."""

    content_types = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    extensions = (".pptx",)

    def extract(self, data: bytes, filename: str) -> ExtractedDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ExtractorUnavailable("python-pptx is not installed") from exc

        try:
            presentation = Presentation(io.BytesIO(data))
            slides: list[str] = []
            for index, slide in enumerate(presentation.slides, start=1):
                lines = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False) and shape.text.strip()
                ]
                if lines:
                    slides.append(f"Slide {index}:\n" + "\n".join(lines))
        except Exception as exc:  # noqa: BLE001 - normalize parser errors
            raise ExtractionError(f"Could not read the presentation: {exc}") from exc

        return ExtractedDocument(
            text="\n\n".join(slides), page_count=len(presentation.slides)
        )


class ImageVisionExtractor(DocumentExtractor):
    """Images via a vision LLM — the OCR path for photographed homework.

    Transcription only: it reads what is on the page and never solves anything,
    so the tutoring prompts stay in full control of pedagogy."""

    content_types = ("image/png", "image/jpeg", "image/jpg", "image/webp", "image/heic")
    extensions = (".png", ".jpg", ".jpeg", ".webp", ".heic")

    _PROMPT = (
        "Transcribe ALL text, math, and questions visible in this image, exactly "
        "as written, preserving question numbers and line breaks. Describe any "
        "diagram or figure briefly in square brackets, e.g. [diagram: a circle "
        "split into 4 equal parts]. Do NOT solve anything, do not add hints, and "
        "do not comment — output only the transcription."
    )

    def extract(self, data: bytes, filename: str) -> ExtractedDocument:
        if not settings.OPENAI_API_KEY:
            raise ExtractorUnavailable("OPENAI_API_KEY is not configured")
        try:
            from openai import OpenAI
        except ImportError as exc:
            raise ExtractorUnavailable("openai is not installed") from exc

        mime = "image/jpeg"
        lowered = filename.lower()
        for ext, candidate in (
            (".png", "image/png"),
            (".webp", "image/webp"),
            (".heic", "image/heic"),
        ):
            if lowered.endswith(ext):
                mime = candidate
                break

        encoded = base64.b64encode(data).decode("ascii")
        start = time.perf_counter()
        try:
            client = OpenAI(
                api_key=settings.OPENAI_API_KEY, timeout=60.0, max_retries=2
            )
            response = client.chat.completions.create(
                model=settings.OPENAI_VISION_MODEL,
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": self._PROMPT},
                            {
                                "type": "image_url",
                                "image_url": {"url": f"data:{mime};base64,{encoded}"},
                            },
                        ],
                    }
                ],
                max_completion_tokens=2000,
            )
            text = (response.choices[0].message.content or "").strip()
        except Exception as exc:  # noqa: BLE001 - normalize SDK/network errors
            raise ExtractionError(f"Could not read the image: {exc}") from exc

        # Metadata only — never log the transcription itself.
        log.info(
            "vision extraction ok model=%s bytes=%d duration_ms=%.1f text_chars=%d",
            settings.OPENAI_VISION_MODEL,
            len(data),
            (time.perf_counter() - start) * 1000,
            len(text),
        )
        if not text:
            raise ExtractionError("No readable text was found in the image.")
        return ExtractedDocument(text=text, page_count=1)


# Ordered by specificity; the first extractor that claims the file wins.
_EXTRACTORS: tuple[DocumentExtractor, ...] = (
    PlainTextExtractor(),
    PdfExtractor(),
    DocxExtractor(),
    PptxExtractor(),
    ImageVisionExtractor(),
)

#: Advertised to the frontend so the file picker and the API agree on one list.
SUPPORTED_EXTENSIONS: tuple[str, ...] = tuple(
    ext for extractor in _EXTRACTORS for ext in extractor.extensions
)


def get_extractor(content_type: str | None, filename: str) -> DocumentExtractor | None:
    """Return the extractor that claims this file, or None if unsupported."""
    for extractor in _EXTRACTORS:
        if extractor.supports(content_type, filename):
            return extractor
    return None
